"""Run the course QA workflow against an explicitly selected practice checkout."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from uuid import uuid4

from agents import make_agents, make_model
from openpyxl import load_workbook
from repository import Repository
from telemetry import NativeTelemetry
from workflow import run_workflow


def read_case(path: Path, case_id: str) -> str:
    if path.suffix.lower() != ".xlsx":
        text = path.read_text(encoding="utf-8-sig")
        if not text.strip():
            raise ValueError("The supplied case file is empty")
        return text
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections = []
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            selected = [row for row in rows if any(str(value).strip() == case_id for value in row)]
            if selected:
                headers = next((row for row in rows if any(value is not None for value in row)), ())
                sections.append(
                    f"## {sheet.title}\n"
                    + json.dumps(
                        {"headers": headers, "rows": selected}, ensure_ascii=False, default=str
                    )
                )
    finally:
        workbook.close()
    if not sections:
        raise ValueError(f"No rows explicitly identify {case_id}; supply its complete text instead")
    return "\n\n".join(sections)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--case-file", type=Path, required=True)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--provider", choices=("anthropic", "openai"), required=True)
    parser.add_argument("--model", required=True)
    parser.add_argument("--api-url", default="http://127.0.0.1:8080")
    parser.add_argument(
        "--otel", action="store_true", help="Export traces using OTEL_* environment settings"
    )
    args = parser.parse_args()
    if not re.fullmatch(r"API-[0-9]+", args.case_id):
        parser.error("--case-id must identify one API case, for example API-1234")

    repository = args.repo.resolve(strict=True)
    case_path = args.case_file.resolve(strict=True)
    case = read_case(case_path, args.case_id)
    output_dir = repository / ".agent-state" / "qa-workflow" / uuid4().hex
    output_dir.mkdir(parents=True)
    adapter = Repository(repository, output_dir, args.case_id, args.api_url)
    telemetry = NativeTelemetry(export=args.otel)
    try:
        agents = make_agents(adapter, lambda: make_model(args.provider, args.model))
        result = run_workflow(case, adapter, agents)
        print(
            json.dumps(
                {"status": result["status"], "report": str(output_dir / "result.json")}, indent=2
            )
        )
        return 0 if result["status"] in {"REVIEWED", "ALREADY_COVERED"} else 1
    finally:
        telemetry.close()


if __name__ == "__main__":
    raise SystemExit(main())
