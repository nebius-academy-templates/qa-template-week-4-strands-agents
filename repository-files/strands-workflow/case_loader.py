"""Load explicitly selected test cases as normalized text."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

TEXT_SUFFIXES = frozenset({".md", ".txt"})
SUPPORTED_SUFFIXES = frozenset({".xlsx", *TEXT_SUFFIXES})
WORKBOOK_SHEETS = ("Case Summary", "Steps")
CASE_ID_HEADER = "Case ID"
REQUIRED_COLUMNS = {
    "Case Summary": (CASE_ID_HEADER, "Title", "Description", "Preconditions"),
    "Steps": (CASE_ID_HEADER, "Actions", "Expected Results"),
}
READINESS_STATUSES: dict[str, Literal["READY", "BLOCKED", "NEEDS_CLARIFICATION"]] = {
    "READY FOR AUTOMATION": "READY",
    "BLOCKED": "BLOCKED",
    "NEEDS_CLARIFICATION": "NEEDS_CLARIFICATION",
}


@dataclass(frozen=True)
class CaseInput:
    case_id: str
    text: str
    readiness_status: Literal["READY", "BLOCKED", "NEEDS_CLARIFICATION"] | None = None


def _as_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _contains_case_id(text: str, case_id: str) -> bool:
    return re.search(rf"(?<![A-Za-z0-9]){re.escape(case_id)}(?![A-Za-z0-9])", text) is not None


def _read_text_case(path: Path, case_ids: list[str]) -> list[CaseInput]:
    if len(case_ids) != 1:
        raise ValueError("A Markdown or text case file requires exactly one case ID")

    case_id = case_ids[0]
    text = path.read_text(encoding="utf-8-sig").strip()
    if not text:
        raise ValueError("The supplied case file is empty")
    if not _contains_case_id(text, case_id):
        raise ValueError(f"The supplied text does not explicitly identify {case_id}")
    return [CaseInput(case_id, text)]


def _rows_by_case(sheet, case_ids: list[str]) -> tuple[list[str], dict[str, list[list[str]]]]:
    rows = [[_as_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    header_rows = [(index, row) for index, row in enumerate(rows) if CASE_ID_HEADER in row]
    if len(header_rows) != 1:
        raise ValueError(
            f"Worksheet {sheet.title!r} must contain exactly one header row with {CASE_ID_HEADER!r}"
        )

    header_index, headers = header_rows[0]
    if headers.count(CASE_ID_HEADER) != 1:
        raise ValueError(
            f"Worksheet {sheet.title!r} must contain exactly one {CASE_ID_HEADER!r} column"
        )
    missing_columns = [column for column in REQUIRED_COLUMNS[sheet.title] if column not in headers]
    if missing_columns:
        names = ", ".join(missing_columns)
        raise ValueError(f"Worksheet {sheet.title!r} is missing required column(s): {names}")

    case_id_column = headers.index(CASE_ID_HEADER)
    selected = {case_id: [] for case_id in case_ids}
    for values in rows[header_index + 1 :]:
        row_case_id = values[case_id_column] if case_id_column < len(values) else ""
        if row_case_id in selected:
            selected[row_case_id].append(values)
    return headers, selected


def _read_workbook_cases(
    path: Path,
    case_ids: list[str],
    require_complete_fields: bool,
) -> list[CaseInput]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = [name for name in WORKBOOK_SHEETS if name not in workbook.sheetnames]
        if missing_sheets:
            names = ", ".join(missing_sheets)
            raise ValueError(f"The workbook is missing required worksheet(s): {names}")

        sections = {case_id: [] for case_id in case_ids}
        readiness_statuses = {case_id: None for case_id in case_ids}
        errors = []
        for sheet_name in WORKBOOK_SHEETS:
            headers, selected = _rows_by_case(workbook[sheet_name], case_ids)
            for case_id in case_ids:
                rows = selected[case_id]
                if len(rows) != 1:
                    count = len(rows)
                    errors.append(
                        f"{sheet_name} must contain exactly one row for {case_id}; found {count}"
                    )
                if len(rows) == 1:
                    row = rows[0]
                    if require_complete_fields:
                        for column in REQUIRED_COLUMNS[sheet_name]:
                            value = row[headers.index(column)]
                            if not value:
                                errors.append(f"{sheet_name} {column!r} is empty for {case_id}")
                    if sheet_name == "Case Summary" and "Automated Test" in headers:
                        marker = row[headers.index("Automated Test")]
                        readiness_statuses[case_id] = READINESS_STATUSES.get(marker)
                sections[case_id].append(
                    f"## {sheet_name}\n"
                    + json.dumps({"headers": headers, "rows": rows}, ensure_ascii=False)
                )

        if errors:
            raise ValueError("; ".join(errors))
        return [
            CaseInput(
                case_id=case_id,
                text="\n\n".join(sections[case_id]),
                readiness_status=readiness_statuses[case_id],
            )
            for case_id in case_ids
        ]
    finally:
        workbook.close()


def read_cases(
    path: Path,
    case_ids: list[str],
    *,
    require_complete_fields: bool = False,
) -> list[CaseInput]:
    """Load every selected case once and preserve the caller's ID order."""
    if not case_ids:
        raise ValueError("At least one case ID is required")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("--case-file must be an XLSX workbook, Markdown file, or text file")
    if suffix in TEXT_SUFFIXES:
        return _read_text_case(path, case_ids)
    return _read_workbook_cases(path, case_ids, require_complete_fields)
